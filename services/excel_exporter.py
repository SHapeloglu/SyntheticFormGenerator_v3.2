from __future__ import annotations

from pathlib import Path
from typing import Any

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter


HEADER_FILL = PatternFill(
    fill_type="solid",
    fgColor="1F4E78",
)

WARNING_FILL = PatternFill(
    fill_type="solid",
    fgColor="FFF2CC",
)


def adjust_column_widths(worksheet) -> None:
    for column_cells in worksheet.columns:
        max_length = 0
        column_letter = get_column_letter(column_cells[0].column)

        for cell in column_cells:
            if cell.value is None:
                continue

            value_length = len(str(cell.value))
            max_length = max(max_length, value_length)

        worksheet.column_dimensions[column_letter].width = min(
            max(max_length + 2, 12),
            40,
        )


def create_excel(
    records: list[dict[str, Any]],
    output_path: Path,
) -> Path:
    if not records:
        raise ValueError("Excel oluşturmak için kayıt bulunamadı.")

    output_path.parent.mkdir(parents=True, exist_ok=True)

    workbook = Workbook()

    worksheet = workbook.active
    worksheet.title = "Sentetik_Veriler"

    headers = list(records[0].keys())
    worksheet.append(headers)

    for record in records:
        worksheet.append(
            [record.get(header, "") for header in headers]
        )

    for cell in worksheet[1]:
        cell.fill = HEADER_FILL
        cell.font = Font(
            bold=True,
            color="FFFFFF",
        )
        cell.alignment = Alignment(
            horizontal="center",
            vertical="center",
            wrap_text=True,
        )

    worksheet.freeze_panes = "A2"
    worksheet.auto_filter.ref = worksheet.dimensions

    for row in worksheet.iter_rows(min_row=2):
        row[1].fill = WARNING_FILL

        for cell in row:
            cell.alignment = Alignment(
                vertical="top",
                wrap_text=True,
            )

    adjust_column_widths(worksheet)

    readme = workbook.create_sheet("README")

    readme["A1"] = "SENTETİK OCR/HTR VERİ SETİ"
    readme["A1"].font = Font(
        bold=True,
        size=16,
        color="FFFFFF",
    )
    readme["A1"].fill = PatternFill(
        fill_type="solid",
        fgColor="C00000",
    )

    readme["A3"] = (
        "Bu dosyadaki kayıtların tamamı yapay olarak üretilmiştir. "
        "Gerçek kişi, çalışan, hasta veya sağlık raporu değildir. "
        "Yalnızca OCR/HTR geliştirme ve test amacıyla kullanılmalıdır. "
        "Kimlik ve SGK numaraları özellikle geçersiz TEST değerleridir. "
        "Hekim adı, imzası ve kaşesi üretilmez."
    )

    readme["A3"].alignment = Alignment(
        wrap_text=True,
        vertical="top",
    )
    readme.column_dimensions["A"].width = 100
    readme.row_dimensions[3].height = 100

    workbook.save(output_path)

    return output_path